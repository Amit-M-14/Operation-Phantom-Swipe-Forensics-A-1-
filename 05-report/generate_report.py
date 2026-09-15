#!/usr/bin/env python3
"""
generate_report.py
--------------------
Generates the 4-6 page Legal-Technical-Report.docx deliverable for
Operation Phantom Swipe from structured content defined in this script.

Run once (from inside 05-report/):
    python generate_report.py

Requires: python-docx  (install with: uv add python-docx)
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ---------- Title page ----------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Operation Phantom Swipe")
run.font.size = Pt(28)
run.bold = True

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Legal-Technical Investigation Report")
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

subtitle2 = doc.add_paragraph()
subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle2.add_run("Cross-Border ATM & Credit Card Fraud Ring")

doc.add_paragraph()
for line in [
    "Assignment 1 — Unit 1: Foundations of Digital Forensics",
    "Prepared by: [Your Name]",
    "Case No.: OPS-2026-014",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(line)

doc.add_page_break()

# ---------- Helper functions ----------
def h1(text):
    doc.add_heading(text, level=1)

def body(text):
    doc.add_paragraph(text)

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

# ---------- 1. Executive Summary ----------
h1("1. Executive Summary")
body(
    "This report documents the early-phase digital forensic investigation into "
    "\u201cOperation Phantom Swipe,\u201d a cross-border ATM skimming and online credit "
    "card fraud ring. The investigation combined physical device seizure (a covert "
    "ATM skimmer), mobile device analysis (a suspect's phone running a disguised "
    "fraud-support application), keyword and metadata-based evidence review, and a "
    "simulated password-recovery exercise on a locked evidence folder. This report "
    "summarizes the crimes identified and their digital footprint, justifies the "
    "evidence-handling procedures followed, discusses the specific challenges of "
    "investigating a cross-border fraud operation, and closes with recommendations "
    "for improving law-enforcement standard operating procedures (SOPs) in similar cases."
)

# ---------- 2. Summary of Crimes and Digital Footprint ----------
h1("2. Summary of Crimes and Digital Footprint")
body(
    "The scenario presents a layered fraud operation spanning hardware and software: "
    "a physical skimming device intercepting card data at an ATM, a mobile application "
    "used to process and potentially clone that data, and coordinated communication "
    "between the suspect and a co-conspirator to move stolen data and launder proceeds "
    "through a money-mule account. Table 1 summarizes the classification of these "
    "crimes and their mapping to applicable law (full justification in "
    "01-classification/cybercrime-taxonomy.md)."
)

doc.add_paragraph("Table 1: Crime-to-Law Mapping")
table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text = "Crime", "Indian Law", "International Framework"

rows_data = [
    ("ATM Skimming", "IT Act \u00a743, \u00a766; IPC \u00a7379, \u00a7420", "Budapest Convention Art. 2, 3"),
    ("Card Cloning", "IT Act \u00a766C; IPC \u00a7468, \u00a7471", "Budapest Convention Art. 7"),
    ("Online/CNP Fraud", "IT Act \u00a766D; IPC \u00a7420", "Budapest Convention Art. 8"),
    ("Organized/Cross-border", "IT Act \u00a766, \u00a743; IPC \u00a7120B, \u00a734", "Budapest Convention Art. 25 (MLA)"),
]
for crime, law, intl in rows_data:
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = crime, law, intl

doc.add_paragraph()
body(
    "The digital footprint recovered during analysis included: two simulated card "
    "track records extracted from the skimmer's flash memory; a sideloaded "
    "application disguised as a calculator utility on the suspect's phone; two "
    "communication log entries referencing a \u201cbatch\u201d of stolen tracks and a "
    "money-mule account (BR-9921); and two GPS location records placing the "
    "suspect's phone at the ATM location both before and after the skimming window. "
    "Individually, each artefact is circumstantial; together, they form a coherent, "
    "corroborated timeline connecting the suspect to the physical installation of "
    "the skimmer, coordination with a co-conspirator, and the financial laundering pathway."
)

# ---------- 3. Evidence Handling Justification ----------
h1("3. Evidence Handling Justification")
body("Evidence handling followed standard digital forensic principles to preserve admissibility:")
bullets([
    "Write-blocking: All extraction from the skimmer's flash memory and the phone was performed through a write-blocked interface, ensuring the acquisition process could not alter source data.",
    "Hashing for integrity: SHA-256 hashes were computed for all six acquired files immediately after extraction (02-evidence-collection/evidence-hashes.sha256) and re-verified before analysis, so that any modification, however small, would be detectable.",
    "Chain of custody: A formal chain-of-custody form documents every seizure, transfer, and access event for both evidence items, with tamper-evident seal numbers logged at each stage.",
    "Isolation of live devices: The suspect's phone was placed in a Faraday bag immediately upon seizure to prevent remote wipe or evidence tampering over a live network connection.",
    "Working on copies only: All keyword search, metadata extraction, and artefact review was performed against verified forensic copies \u2014 never the sealed original media.",
])
body(
    "These measures collectively address the two central evidentiary concerns in any "
    "digital forensic proceeding: authenticity (is this really the data recovered "
    "from the seized device?) and integrity (has the data been altered since acquisition?)."
)

# ---------- 4. Cross-Border Challenges ----------
h1("4. Challenges in Cross-Border Investigation and Cooperation")
body(
    "Because the fraud ring's digital footprint spans jurisdictions \u2014 skimmed "
    "data captured domestically, potentially exfiltrated or monetized through "
    "offshore accounts, and a co-conspirator possibly operating from another "
    "country \u2014 several structural challenges arise:"
)
bullets([
    "Jurisdictional authority: Indian law enforcement's search and seizure powers under the IT Act and CrPC/BNSS do not extend automatically to servers, accounts, or suspects located abroad, requiring formal international cooperation mechanisms.",
    "Mutual Legal Assistance Treaties (MLATs): Obtaining evidence held by foreign service providers or banks typically requires an MLAT request, which can take months \u2014 a significant mismatch against the speed at which digital evidence can be deleted or overwritten.",
    "Non-signatory status: India is not a signatory to the Budapest Convention, which streamlines cooperation among member states; this limits the fast-track cooperation channels available compared to signatory-to-signatory cases.",
    "Data localization and privacy law conflicts: Differing data protection regimes (e.g., GDPR in the EU vs. India's DPDP Act) can complicate or delay requests for subscriber or transaction data from foreign platforms.",
    "Attribution difficulty: VPNs, offshore hosting, and the use of a money-mule account (rather than the principal's own account) are specifically designed to frustrate attribution back to the actual beneficiaries of the fraud.",
    "Evidence format and standards mismatch: Forensic evidence collected under one jurisdiction's procedural standards may face admissibility challenges if presented in another jurisdiction's courts without matching certification/authentication requirements.",
])

# ---------- 5. Recommendations ----------
h1("5. Recommendations for Law Enforcement SOP Improvements")
bullets([
    "Establish a fast-track digital-evidence preservation request (a \u201cquick-freeze\u201d mechanism) that can be issued to foreign providers ahead of a full MLAT request, to prevent data deletion during the months-long MLAT process.",
    "Build standing points of contact with major payment networks and banks for real-time flagging of skimming-pattern transactions, enabling earlier detection before a ring can escalate to cross-border cash-out.",
    "Mandate hash-verification checkpoints at every custody transfer (not just at acquisition), with automated logging, to strengthen the evidentiary chain for cases likely to involve multi-agency or multi-country handoffs.",
    "Adopt a standardized cross-border evidence packaging format (metadata + hash manifest + custody log bundled together) so that evidence collected domestically is readily admissible or at least readily verifiable when shared with foreign counterparts.",
    "Invest in weak-password-aware forensic tooling (dictionary and rule-based cracking, as demonstrated in Section 4 of the technical deliverables) as a standard first step before resorting to lengthy brute-force or compelled-decryption processes, reserving compelled decryption for cases where technical recovery fails.",
    "Formalize training on money-mule account identification and rapid asset-freeze procedures, since mule accounts are often the fastest-closing window for recovering or halting the flow of fraudulent proceeds.",
])

# ---------- 6. Conclusion ----------
h1("6. Conclusion")
body(
    "The Operation Phantom Swipe simulation demonstrates how a combination of "
    "physical device forensics, mobile artefact analysis, and disciplined evidence "
    "handling can reconstruct a coherent fraud timeline even from a small number of "
    "artefacts. The greatest practical obstacles are not technical but procedural "
    "and jurisdictional \u2014 the speed mismatch between how quickly digital evidence "
    "can vanish and how slowly cross-border legal cooperation currently moves. "
    "Closing that gap through faster preservation mechanisms and standardized "
    "evidence handling is the single highest-leverage SOP improvement identified "
    "in this investigation."
)

# ---------- Appendix ----------
h1("Appendix: Deliverables Cross-Reference")
bullets([
    "Cybercrime taxonomy & legal mapping: 01-classification/cybercrime-taxonomy.md",
    "Chain of custody & acquisition log: 02-evidence-collection/",
    "Evidence hashes (SHA-256, 6 files): 02-evidence-collection/evidence-hashes.sha256",
    "Search strategy, evidence log, extracted artefacts: 03-media-analysis/",
    "Cryptography simulation & ethics discussion: 04-cryptography/",
    "Tools (hash generator, string search, metadata extractor, password cracker): tools/ and 04-cryptography/crack_password.py",
])

doc.save("Legal-Technical-Report.docx")
print("Created Legal-Technical-Report.docx")
